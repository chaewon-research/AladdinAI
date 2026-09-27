"""Generate synthetic, non-confidential WRT preservation fixtures (requires python-docx, Pillow)."""
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image

HERE = Path(__file__).parent / 'fixtures'
HERE.mkdir(exist_ok=True)

def img(fmt, color):
    b = BytesIO()
    im = Image.new('RGB', (40, 25), color)
    im.save(b, format=fmt)
    b.seek(0)
    return b

def write(name, d):
    d.save(HERE / name)
    print('GENERATED', name)

# A straightforward control containing mixed body formatting.
d=Document(); d.add_heading('Control',1)
p=d.add_paragraph(); p.add_run('Normal '); p.add_run('Bold').bold=True; p.add_run(' '); p.add_run('Italic').italic=True; p.add_run(' Unchanged.')
d.add_paragraph('This entire paragraph should stay intact.')
write('text_mixed.docx', d)

# Simple table; a supported WRT representation.
d=Document(); d.add_heading('Table control',1)
t=d.add_table(rows=2, cols=2); t.cell(0,0).text='Header'; t.cell(0,1).text='Other'; t.cell(1,0).text='First'; t.cell(1,1).text='Keep'
write('table_plain.docx', d)

# Bold within a table cell (common, relevant for semantic preservation).
d=Document(); d.add_heading('Formatted table',1)
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Header'; t.cell(0,1).text='Other'; p=t.cell(1,0).paragraphs[0]; p.add_run('BoldCell').bold=True; t.cell(1,1).text='Keep'
write('table_formatted.docx', d)

# A legitimate cell contains the WRT table delimiter.
d=Document(); d.add_heading('Pipe in table',1)
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Header'; t.cell(0,1).text='Other'; t.cell(1,0).text='A | B'; t.cell(1,1).text='Keep'
write('table_pipe.docx', d)

# Inline picture appears between two text runs in ONE paragraph.
d=Document(); d.add_heading('Inline image',1)
p=d.add_paragraph(); p.add_run('Before picture '); p.add_run().add_picture(img('PNG','red'), width=Inches(.4)); p.add_run(' after picture')
write('inline_png.docx', d)

# The original picture is JPEG; round-tripped OOXML should identify it correctly.
d=Document(); d.add_heading('JPEG image',1)
p=d.add_paragraph('JPEG photo: '); p.add_run().add_picture(img('JPEG','blue'), width=Inches(.4))
write('jpeg_image.docx', d)

# One targeted change, with text/table/image outside its scope.
d=Document(); d.add_heading('Edit scope',1); d.add_paragraph('Target: DRAFT'); d.add_paragraph('Untouched paragraph.')
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Field'; t.cell(0,1).text='Value'; t.cell(1,0).text='Item'; t.cell(1,1).text='Keep'
p=d.add_paragraph(); p.add_run('Before picture '); p.add_run().add_picture(img('PNG','green'), width=Inches(.4)); p.add_run(' after picture')
write('targeted_edit.docx', d)
